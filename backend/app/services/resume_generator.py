"""简历：必要问题生成 + 个人简介 + .docx 生成。"""
import json
import re
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import llm_client

PRIMARY = RGBColor(0x1F, 0x4E, 0x79)
_SYSTEM = "你是一名资深技术招聘与简历专家。"


# ---------- AI 生成部分 ----------
async def generate_resume_questions(
    jd_summary: dict,
    user_profile: dict,
    base_url: str,
    api_key: str,
    model: str,
) -> list:
    prompt = (
        "请结合岗位 JD 与求职者现有档案，判断简历中还缺失或不确定的信息，"
        "设计 3~5 个必要问题（用于补齐定向简历所需关键素材）。\n"
        "只返回 JSON 数组（不要 markdown 代码块），每项结构：\n"
        '{"id":"q1","type":"choice|multi|text","question":"问题内容","options":["选项"],"required":true}\n\n'
        f"JD 解析结果：{json.dumps(jd_summary, ensure_ascii=False)}\n"
        f"求职者档案：{json.dumps(user_profile, ensure_ascii=False)}\n\n"
        "问题应聚焦：意向岗位确认、项目真实性与量化数据、JD 匹配的技术点等。"
    )
    raw = await llm_client.chat(
        base_url, api_key, model,
        [{"role": "system", "content": _SYSTEM}, {"role": "user", "content": prompt}],
        temperature=0.4, json_mode=True,
    )
    return _parse_questions(raw)


def _parse_questions(raw: str) -> list:
    try:
        data = json.loads(raw)
    except Exception:
        return []
    if isinstance(data, dict):
        data = data.get("questions") or data.get("data") or []
    if not isinstance(data, list):
        return []
    out = []
    for i, item in enumerate(data):
        if not isinstance(item, dict):
            continue
        out.append({
            "id": str(item.get("id") or f"q{i + 1}"),
            "type": item.get("type") or "text",
            "question": item.get("question") or "",
            "options": item.get("options") or [],
            "required": bool(item.get("required", True)),
        })
    return out


async def generate_personal_summary(
    jd_summary: dict,
    user_profile: dict,
    base_url: str,
    api_key: str,
    model: str,
) -> str:
    prompt = (
        "请为求职者写一段 2~3 行的个人简介（电梯陈述），3 秒讲清「我是谁 + 为什么匹配该岗位」，"
        "包含技术栈速览。要求口语化、简洁、有针对性。\n"
        "关键量化成果与核心技术栈用 ** 包裹（例如 **85%**、**FastAPI**），用于加粗高亮；"
        "除此之外不要使用其他 markdown 语法。\n\n"
        f"JD 解析结果：{json.dumps(jd_summary, ensure_ascii=False)}\n"
        f"求职者档案：{json.dumps(user_profile, ensure_ascii=False)}\n\n"
        "只输出简介正文，不要标题。"
    )
    return await llm_client.chat(
        base_url, api_key, model,
        [{"role": "system", "content": _SYSTEM}, {"role": "user", "content": prompt}],
        temperature=0.6,
    )


async def curate_skills(
    jd_summary: dict,
    skills: list,
    base_url: str,
    api_key: str,
    model: str,
    max_count: int = 10,
) -> list:
    """技能过多时，由 AI 挑选与 JD 最贴近的技能并按相关度降序排列。"""
    if not skills:
        return []
    items = [
        {"category": s["category"], "name": s["name"], "proficiency": s["proficiency"]}
        for s in skills
    ]
    prompt = (
        f"以下是求职者的专业技能列表（共 {len(items)} 项）。请从中挑选与岗位最相关、"
        f"最值得写入简历的最多 {max_count} 项，并按与 JD 的匹配度从高到低排序。\n"
        "只返回 JSON 数组（不要 markdown 代码块、不要额外文字），数组元素为技能名（字符串），"
        "必须原样取自下面的列表，不要改动名称。\n\n"
        f"岗位 JD 解析结果：{json.dumps(jd_summary, ensure_ascii=False)}\n"
        f"技能列表：{json.dumps([s['name'] for s in items], ensure_ascii=False)}"
    )
    raw = await llm_client.chat(
        base_url, api_key, model,
        [{"role": "system", "content": _SYSTEM}, {"role": "user", "content": prompt}],
        temperature=0.2, json_mode=True,
    )
    try:
        order = json.loads(raw)
    except Exception:
        order = []
    if isinstance(order, dict):
        order = order.get("skills") or order.get("data") or []
    if not isinstance(order, list):
        return []
    by_name = {s["name"]: s for s in items}
    result = []
    for name in order:
        if not isinstance(name, str):
            continue
        if name in by_name:
            result.append(by_name.pop(name))
    return result


# ---------- DOCX 构建 ----------
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)
GRAY = RGBColor(0x8A, 0x8A, 0x8A)
DARK = RGBColor(0x2B, 0x2B, 0x2B)

_EMPH_RE = re.compile(r"(\*\*.+?\*\*)")


def _set_run(run, size=10.5, bold=False, color=DARK, italic=False):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def _border(p, color="1F4E79", sz="6", side="bottom", space="1"):
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), space)
    el.set(qn("w:color"), color)
    pbdr.append(el)
    ppr.append(pbdr)


def _add_para(container):
    """取容器中的可写段落：表格单元格首个空段落复用，否则新增。Document 与单元格通用。"""
    ps = container.paragraphs
    if len(ps) == 1 and not ps[0].runs:
        return ps[0]
    return container.add_paragraph()


def _rule(container, color="1F4E79", sz="10"):
    p = _add_para(container)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    _border(p, color=color, sz=sz, side="bottom", space="1")
    return p


def _add_rich_runs(p, text, size=10, color=DARK, bold=False):
    """渲染带强调标记的文本：**内容** 渲染为蓝色加粗，其余正常。"""
    for seg in _EMPH_RE.split(text or ""):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            r = p.add_run(seg[2:-2])
            _set_run(r, size=size, bold=True, color=ACCENT)
        else:
            r = p.add_run(seg)
            _set_run(r, size=size, bold=bold, color=color)


def _heading(container, text):
    p = _add_para(container)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r0 = p.add_run("■ ")
    _set_run(r0, size=9, bold=True, color=ACCENT)
    r = p.add_run(text)
    _set_run(r, size=12.5, bold=True, color=PRIMARY)
    _border(p, color="B7C9E2", sz="6", side="bottom", space="2")
    return p


def _entry_head(container, left, right, left_bold=True, tab_cm=17.3):
    p = _add_para(container)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(tab_cm), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left)
    _set_run(r, size=11, bold=left_bold, color=DARK)
    if right:
        r2 = p.add_run("\t" + right)
        _set_run(r2, size=9, color=GRAY)
    return p


def _bullet(container, text, size=10, color=DARK):
    p = _add_para(container)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run("· ")
    _set_run(r, size=size, bold=True, color=ACCENT)
    _add_rich_runs(p, text, size=size, color=color)
    return p


def _tech_line(container, text):
    p = _add_para(container)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Cm(0.4)
    r = p.add_run("技术栈  ")
    _set_run(r, size=9.5, bold=True, color=ACCENT)
    r2 = p.add_run(text)
    _set_run(r2, size=9.5, color=GRAY)
    return p


_PROF_LEVEL = {"熟练": 90, "精通": 95, "掌握": 80, "熟悉": 70, "了解": 50}


def _proficiency_level(label, default=60):
    return _PROF_LEVEL.get((label or "").strip(), default)


# 进度条尺寸：3.0cm × 0.16cm（EMU）
_BAR_W = 1080000
_BAR_H = 58000
_bar_id = iter(range(10001, 99999))


def _skill_bar(p, level, fill="1F6FB2", bg="E4E9F0"):
    """用 DrawingML 圆角矩形 + 双色硬渐变渲染图形化进度条（纯 XML，无图片）。

    level: 0~100 的整数百分比；fill 为已填充段颜色，bg 为底色。
    """
    level = max(0, min(100, int(level)))
    pos = level * 1000  # 渐变位置（千分比）
    sid = next(_bar_id)
    xml = (
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
        f'<wp:extent cx="{_BAR_W}" cy="{_BAR_H}"/>'
        f'<wp:docPr id="{sid}" name="skillbar{sid}"/>'
        '<wp:cNvGraphicFramePr/>'
        '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        '<wps:wsp>'
        f'<wps:cNvSpPr id="{sid}" name="skillbar{sid}"/>'
        '<wps:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{_BAR_W}" cy="{_BAR_H}"/></a:xfrm>'
        '<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 50000"/></a:avLst></a:prstGeom>'
        '<a:gradFill><a:gsLst>'
        f'<a:gs pos="0"><a:srgbClr val="{fill}"/></a:gs>'
        f'<a:gs pos="{pos}"><a:srgbClr val="{fill}"/></a:gs>'
        f'<a:gs pos="{pos}"><a:srgbClr val="{bg}"/></a:gs>'
        f'<a:gs pos="100000"><a:srgbClr val="{bg}"/></a:gs>'
        '</a:gsLst><a:lin ang="0" scaled="0"/></a:gradFill>'
        '<a:ln><a:noFill/></a:ln>'
        '</wps:spPr>'
        '<wps:bodyPr/>'
        '</wps:wsp>'
        '</a:graphicData></a:graphic>'
        '</wp:inline></w:drawing></w:r>'
    )
    p._p.append(parse_xml(xml))


def _skill_cell(p, name, proficiency):
    """技能条目：名称 + 图形化进度条（XML 渲染）+ 熟练度描述词。"""
    level = _proficiency_level(proficiency)
    r = p.add_run(name)
    _set_run(r, size=9.5, bold=True, color=DARK)
    r = p.add_run(" ")
    _set_run(r, size=9)
    _skill_bar(p, level)
    if proficiency:
        r = p.add_run(" " + proficiency)
        _set_run(r, size=8.5, color=GRAY)


def _cell_width(name, proficiency):
    """估算技能条目显示宽度（以全角字符为单位），用于判断能否并列一行。"""
    w = sum(1.0 if ord(ch) > 127 else 0.55 for ch in (name or ""))
    w += 9.0  # 图形化进度条固定宽 3.0cm
    if proficiency:
        w += len(proficiency) * 1.0
    return w


_SKILL_TAB = Cm(8.8)
_SKILL_HALF = 24.0  # 半行可容纳的全角字符宽度


def _skill_row(container, skills):
    """技能条目排版：空间允许时每行并列两个（用制表位对齐），过长则独占一行。"""
    i = 0
    n = len(skills)
    while i < n:
        name, prof = skills[i]
        if (i + 1 < n and _cell_width(name, prof) <= _SKILL_HALF
                and _cell_width(*skills[i + 1]) <= _SKILL_HALF):
            pair = [skills[i], skills[i + 1]]
            i += 2
        else:
            pair = [skills[i]]
            i += 1
        p = _add_para(container)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.tab_stops.add_tab_stop(_SKILL_TAB, WD_TAB_ALIGNMENT.LEFT)
        _skill_cell(p, pair[0][0], pair[0][1])
        if len(pair) == 2:
            r = p.add_run("\t")
            _set_run(r, size=9)
            _skill_cell(p, pair[1][0], pair[1][1])


def _enabled(sections, key):
    return not sections or key in sections


def _merge_sections(user, fabricated):
    """合并杜撰经历到项目/实习/工作/其他列表，返回四元组。"""
    fake_projects = [SimpleNamespace(**x) for x in (fabricated or []) if x.get("type") == "project"]
    fake_internships = [
        SimpleNamespace(
            company=x.get("title", ""), position=x.get("role", ""),
            start_date=x.get("start_date", ""), end_date=x.get("end_date", ""),
            description=x.get("description", ""), achievement=x.get("result", ""),
        )
        for x in (fabricated or []) if x.get("type") == "internship"
    ]
    fake_works = [
        SimpleNamespace(
            company=x.get("title", ""), position=x.get("role", ""),
            start_date=x.get("start_date", ""), end_date=x.get("end_date", ""),
            description=x.get("description", ""), achievement=x.get("result", ""),
        )
        for x in (fabricated or []) if x.get("type") == "work"
    ]
    fake_others = [
        SimpleNamespace(
            company=x.get("title", ""), position=x.get("role", ""),
            start_date=x.get("start_date", ""), end_date=x.get("end_date", ""),
            description=x.get("description", ""), achievement=x.get("result", ""),
        )
        for x in (fabricated or []) if x.get("type") == "other"
    ]
    projects = list(user.projects) + fake_projects
    internships = list(user.internships) + fake_internships
    works = fake_works
    others = fake_others
    return projects, internships, works, others


def _new_doc():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    return doc


def _save_docx(doc, user, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    fname = f"resume_{user.id}_{stamp}.docx"
    out_path = out_dir / fname
    doc.save(str(out_path))
    return str(out_path)


def build_resume_docx(user, jd_summary, personal_summary, self_evaluation, out_dir: Path,
                      target_position: str = "", sections=None, fabricated=None,
                      skills_curated=None, expected_salary: str = "", include_salary: bool = True,
                      template: str = "classic") -> str:
    spec = TEMPLATES.get(template)
    if spec is None:
        raise ValueError(f"未知简历模板：{template}")
    return spec["render"](user, jd_summary, personal_summary, self_evaluation, out_dir,
                          target_position, sections, fabricated, skills_curated,
                          expected_salary, include_salary)


def _render_classic(user, jd_summary, personal_summary, self_evaluation, out_dir: Path,
                    target_position: str = "", sections=None, fabricated=None,
                    skills_curated=None, expected_salary: str = "", include_salary: bool = True) -> str:
    # sections 为空/None 表示包含全部模块；否则仅包含勾选模块
    def enabled(key: str) -> bool:
        return _enabled(sections, key)

    projects, internships, works, others = _merge_sections(user, fabricated)

    doc = _new_doc()

    # ===== 头部 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(user.name or "姓名")
    _set_run(r, size=20, bold=True, color=PRIMARY)

    pos = target_position or user.target_position or jd_summary.get("position") or "软件工程师"
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(pos)
    _set_run(r, size=12, bold=True, color=ACCENT)

    if include_salary and expected_salary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("期望薪资  ")
        _set_run(r, size=10.5, bold=True, color=GRAY)
        r2 = p.add_run(expected_salary)
        _set_run(r2, size=11, bold=True, color=ACCENT)

    items = [("电话", user.phone), ("邮箱", user.email), ("城市", user.location)]
    items = [(k, v) for k, v in items if v]
    if items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        for i, (k, v) in enumerate(items):
            if i:
                r = p.add_run("    ")
                _set_run(r, size=9, color=GRAY)
            r = p.add_run(k + " ")
            _set_run(r, size=9, color=GRAY)
            r2 = p.add_run(v)
            _set_run(r2, size=9.5, color=DARK)
    _rule(doc, color="1F4E79", sz="12")

    # 个人简介
    if enabled("summary") and personal_summary:
        _heading(doc, "个人简介")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, personal_summary, size=10)

    # 教育背景
    if enabled("education") and user.educations:
        _heading(doc, "教育背景")
        for edu in user.educations:
            title = " | ".join(x for x in (edu.school, edu.major, edu.degree) if x)
            time = f"{edu.start_date} - {edu.end_date}".strip(" -")
            _entry_head(doc, title, time, left_bold=True)
            if edu.courses:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run("主修课程  ")
                _set_run(r, size=9, bold=True, color=GRAY)
                _add_rich_runs(p, edu.courses, size=9, color=DARK)
            extra = " · ".join(x for x in (f"GPA {edu.gpa}" if edu.gpa else "", f"排名 {edu.rank}" if edu.rank else "") if x)
            if extra:
                _bullet(doc, extra, size=9)

    # 专业技能：进度条展示熟练度，空间允许时至多两个条目一行（制表位对齐）
    if enabled("skills") and (skills_curated or user.skills):
        _heading(doc, "专业技能")
        if skills_curated:
            items = [
                (s.get("name", "") if isinstance(s, dict) else s.name,
                 s.get("proficiency", "") if isinstance(s, dict) else s.proficiency)
                for s in skills_curated
            ]
            _skill_row(doc, items)
        else:
            grouped = {}
            for s in user.skills:
                grouped.setdefault(s.category or "其他", []).append(s)
            for cat, sks in grouped.items():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(cat + "  ")
                _set_run(r, size=10.5, bold=True, color=PRIMARY)
                _skill_row(doc, [(s.name, s.proficiency) for s in sks])

    # 实习经历
    if enabled("internships") and internships:
        _heading(doc, "实习经历")
        for it in internships:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(doc, title, time, left_bold=True)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 工作经历（杜撰）
    if enabled("work") and works:
        _heading(doc, "工作经历")
        for it in works:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(doc, title, time, left_bold=True)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 其他经历（杜撰）
    if enabled("other") and others:
        _heading(doc, "其他经历")
        for it in others:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(doc, title, time, left_bold=True)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 项目经历
    if enabled("projects") and projects:
        _heading(doc, "项目经历")
        for proj in projects:
            time = f"{proj.start_date} - {proj.end_date}".strip(" -")
            _entry_head(doc, proj.title, time, left_bold=True)
            if proj.tech_stack:
                _tech_line(doc, proj.tech_stack)
            if proj.description:
                _bullet(doc, proj.description)
            if proj.result:
                _bullet(doc, proj.result)

    # 证书 / 获奖
    if enabled("certificates") and user.certificates:
        _heading(doc, "证书 / 获奖")
        for c in user.certificates:
            text = c.name + (f"（{c.date}）" if c.date else "")
            _bullet(doc, text)

    # 自我评价（放在简历最下面一栏）
    if enabled("self_evaluation") and self_evaluation:
        _heading(doc, "自我评价")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, self_evaluation, size=10)

    return _save_docx(doc, user, out_dir)


def _cell_shading(cell, color):
    """给表格单元格设置浅色底纹。"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def _no_table_borders(table):
    """去掉表格所有边框，避免双栏版式出现表格线。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _side_heading(container, text):
    p = _add_para(container)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    _set_run(r, size=11, bold=True, color=PRIMARY)
    _border(p, color="B7C9E2", sz="6", side="bottom", space="2")
    return p


def _side_info(container, label, value):
    if not value:
        return
    p = _add_para(container)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label + "  ")
    _set_run(r, size=9, bold=True, color=GRAY)
    r2 = p.add_run(value)
    _set_run(r2, size=9.5, color=DARK)


def _skill_lines(container, skills):
    """侧边栏技能：每行一个技能，窄栏内不并列。"""
    for name, prof in skills:
        p = _add_para(container)
        p.paragraph_format.space_after = Pt(3)
        _skill_cell(p, name, prof)


def _skill_groups(user, skills_curated):
    """返回 [(分类, [(技能名, 熟练度), ...])]；curated 时归为单一分类。"""
    if skills_curated:
        items = [
            (s.get("name", "") if isinstance(s, dict) else s.name,
             s.get("proficiency", "") if isinstance(s, dict) else s.proficiency)
            for s in skills_curated
        ]
        return [("专业技能", items)]
    grouped = {}
    for s in user.skills:
        grouped.setdefault(s.category or "其他", []).append((s.name, s.proficiency))
    return [(cat, grouped[cat]) for cat in grouped]


def _render_split(user, jd_summary, personal_summary, self_evaluation, out_dir: Path,
                  target_position: str = "", sections=None, fabricated=None,
                  skills_curated=None, expected_salary: str = "", include_salary: bool = True) -> str:
    def enabled(key: str) -> bool:
        return _enabled(sections, key)

    projects, internships, works, others = _merge_sections(user, fabricated)

    doc = _new_doc()
    # 双栏：左列侧边栏，右列主栏，隐藏边框
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    _no_table_borders(table)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.width = Cm(6.0)
    right.width = Cm(11.4)
    _cell_shading(left, "EAF1F8")

    # ---- 左侧边栏 ----
    pos = target_position or user.target_position or jd_summary.get("position") or "软件工程师"
    p = _add_para(left)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(user.name or "姓名")
    _set_run(r, size=16, bold=True, color=PRIMARY)
    p = _add_para(left)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(pos)
    _set_run(r, size=10, bold=True, color=ACCENT)

    _side_info(left, "电话", user.phone)
    _side_info(left, "邮箱", user.email)
    _side_info(left, "城市", user.location)
    if include_salary and expected_salary:
        _side_info(left, "期望薪资", expected_salary)

    skill_groups = _skill_groups(user, skills_curated)
    if enabled("skills") and skill_groups:
        _side_heading(left, "专业技能")
        for cat, sks in skill_groups:
            if len(skill_groups) > 1:
                p = _add_para(left)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(cat)
                _set_run(r, size=9, bold=True, color=ACCENT)
            _skill_lines(left, sks)

    if enabled("certificates") and user.certificates:
        _side_heading(left, "证书 / 获奖")
        for c in user.certificates:
            text = c.name + (f"（{c.date}）" if c.date else "")
            p = _add_para(left)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run("· " + text)
            _set_run(r, size=9, color=DARK)

    if getattr(user, "hobbies", ""):
        _side_heading(left, "兴趣爱好")
        p = _add_para(left)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(user.hobbies)
        _set_run(r, size=9, color=DARK)

    # ---- 右侧主栏 ----
    if enabled("summary") and personal_summary:
        _heading(right, "个人简介")
        p = _add_para(right)
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, personal_summary, size=10)

    if enabled("education") and user.educations:
        _heading(right, "教育背景")
        for edu in user.educations:
            title = " | ".join(x for x in (edu.school, edu.major, edu.degree) if x)
            time = f"{edu.start_date} - {edu.end_date}".strip(" -")
            _entry_head(right, title, time, left_bold=True, tab_cm=10.6)
            if edu.courses:
                p = _add_para(right)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run("主修课程  ")
                _set_run(r, size=9, bold=True, color=GRAY)
                _add_rich_runs(p, edu.courses, size=9, color=DARK)
            extra = " · ".join(x for x in (f"GPA {edu.gpa}" if edu.gpa else "", f"排名 {edu.rank}" if edu.rank else "") if x)
            if extra:
                _bullet(right, extra, size=9)

    if enabled("internships") and internships:
        _heading(right, "实习经历")
        for it in internships:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(right, title, time, left_bold=True, tab_cm=10.6)
            if it.description:
                _bullet(right, it.description)
            if it.achievement:
                _bullet(right, it.achievement)

    if enabled("work") and works:
        _heading(right, "工作经历")
        for it in works:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(right, title, time, left_bold=True, tab_cm=10.6)
            if it.description:
                _bullet(right, it.description)
            if it.achievement:
                _bullet(right, it.achievement)

    if enabled("other") and others:
        _heading(right, "其他经历")
        for it in others:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _entry_head(right, title, time, left_bold=True, tab_cm=10.6)
            if it.description:
                _bullet(right, it.description)
            if it.achievement:
                _bullet(right, it.achievement)

    if enabled("projects") and projects:
        _heading(right, "项目经历")
        for proj in projects:
            time = f"{proj.start_date} - {proj.end_date}".strip(" -")
            _entry_head(right, proj.title, time, left_bold=True, tab_cm=10.6)
            if proj.tech_stack:
                _tech_line(right, proj.tech_stack)
            if proj.description:
                _bullet(right, proj.description)
            if proj.result:
                _bullet(right, proj.result)

    if enabled("self_evaluation") and self_evaluation:
        _heading(right, "自我评价")
        p = _add_para(right)
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, self_evaluation, size=10)

    return _save_docx(doc, user, out_dir)


# ---------- 模板库 ----------
# 新增模板时在此登记：key 为唯一标识（持久化到 Resume.template），label 用于前端展示，
# render 为渲染函数。渲染函数统一签名：
#   (user, jd_summary, personal_summary, self_evaluation, out_dir,
#    target_position, sections, fabricated, skills_curated,
#    expected_salary, include_salary) -> str(输出文件路径)
TEMPLATES = {
    "classic": {"label": "简洁单栏", "render": _render_classic},
    "split": {"label": "双栏", "render": _render_split},
}

# 独立模板文件夹（backend/app/services/templates/）中的模板在此登记；
# 置于此处在末尾导入，避免与模板模块之间的循环依赖。
from .templates import blue_classic, ios_engineer  # noqa: E402

TEMPLATES.update(ios_engineer.TEMPLATE)
TEMPLATES.update(blue_classic.TEMPLATE)


def available_templates():
    """返回可用模板列表，供前端渲染选择器与路由校验使用。"""
    return [{"key": key, "label": spec["label"]} for key, spec in TEMPLATES.items()]


def has_template(key: str) -> bool:
    return key in TEMPLATES