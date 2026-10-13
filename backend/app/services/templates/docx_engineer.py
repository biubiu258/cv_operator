"""工程师简历模板（基于真实 docx 模板的填充式渲染）。

与 ``blue_classic`` / ``ios_engineer`` 等「代码逐行拼 docx」的方案不同，
本模板直接打开 ``assets/engineer.docx`` 这份已设计好的 Word 简历模板，
按段落槽位把简历数据填进去，保留模板原有的三栏版式、图标、浮动装饰元素。

实现要点：
1. ``_set_para_text`` 用 lxml 直接操作 ``w:r``，解决模板中文字被拆成多个 run
   或被 ``<w:hyperlink>`` 包裹导致 ``python-docx`` 的 ``paragraph.runs`` 拿不到的问题。
2. 槽位按「段落索引」定位；经历/技能条数超过模板固定槽位时截断，不足时清空。
3. 浮动竖排文本框内的文字不在 ``doc.paragraphs`` 顶层，首版不替换。
"""
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from docx.oxml.ns import qn

from ..prompting import get_length_budget
from ..resume_generator import _enabled, _merge_sections, _save_docx

TEMPLATE_PATH = Path(__file__).parent / "assets" / "engineer.docx"

# 模板固定槽位（以 doc.paragraphs 的顶层 w:p 顺序为准）
SLOT_NAME = 5
SLOT_TARGET = 6
SLOT_AGE = 7
SLOT_ADDR_PHONE = 8
SLOT_EMAIL = 9
SLOT_SELF_1 = 17
SLOT_SELF_2 = 18
SLOT_SKILL_1 = 19
SLOT_SKILL_2 = 24
SLOT_PROJ_TITLE = 26
SLOT_PROJ_TIME_ROLE = 27
SLOT_PROJ_DESC_1 = 29
SLOT_PROJ_DESC_2 = 30
SLOT_FOREIGN = 31
SLOT_EDU_SCHOOL = 38
SLOT_WORK1_COMPANY = 40
SLOT_WORK1_POSITION = 41
SLOT_WORK1_DESC = 42
SLOT_WORK2_TITLE = 43
SLOT_WORK2_POSITION = 44
SLOT_WORK2_DESC = 45
SLOT_WORK1_TIME = 47
SLOT_WORK2_TIME = 52
SLOT_EDU_TIME = 53
SLOT_EDU_MAJOR = 54
SLOT_EDU_COURSE_1 = 55
SLOT_EDU_COURSE_2 = 56
SLOT_EDU_COURSE_3 = 57
SLOT_PROJ2_DESC = 59
SLOT_PROJ3_TITLE = 60
SLOT_WORK3_POSITION = 66
SLOT_WORK3_DESC = 67
SLOT_WORK3_TIME = 68
SLOT_SPEC_1 = 70
SLOT_WORK3_DESC2 = 72
SLOT_SPEC_2 = 75
SLOT_SPEC_3 = 76
SLOT_SPEC_4 = 77
SLOT_HOBBY_1 = 79
SLOT_HOBBY_2 = 80
SLOT_SPEC_5 = 81


def _set_para_text(paragraph, text: str) -> None:
    """把段落整段文字替换为 ``text``，保留首 run 的 rPr 格式。

    兼容两种情况：
    - 文字被拆成多个 run：保留第一个 run，删去其余 run。
    - 文字包在 ``<w:hyperlink>`` 里（``paragraph.runs`` 为空）：找到段落内
      第一个 ``w:r`` 元素同样处理。
    """
    p = paragraph._p
    runs = p.findall(qn("w:r"))
    # 若直接子节点没有 run，再在全段落内找（含 hyperlink 内的 run）
    if not runs:
        runs = p.iter(qn("w:r"))
        runs = list(runs)
    if runs:
        first = runs[0]
        # 清空首 run 原文字，写入新文字
        for t in first.findall(qn("w:t")):
            first.remove(t)
        t_el = first.makeelement(qn("w:t"), {})
        t_el.text = text
        # 保留空格
        t_el.set(qn("xml:space"), "preserve")
        first.append(t_el)
        # 删除其余 run
        for r in runs[1:]:
            parent = r.getparent()
            if parent is not None:
                parent.remove(r)
    else:
        # 段落无 run，新增一个
        r = p.makeelement(qn("w:r"), {})
        t_el = r.makeelement(qn("w:t"), {})
        t_el.text = text
        t_el.set(qn("xml:space"), "preserve")
        r.append(t_el)
        p.append(r)


def _clear(paragraph) -> None:
    """清空段落文字（保留段落本身与格式）。"""
    _set_para_text(paragraph, "")


def _split_text(text: str, max_chars: int = 60) -> list:
    """把一段长文本按句号/分号切成至多两段子文本，适配模板双行槽位。"""
    if not text:
        return ["", ""]
    text = text.strip()
    if len(text) <= max_chars:
        return [text, ""]
    # 优先在句号/分号处切
    cut = -1
    for sep in ("。", "；", ";"):
        idx = text.find(sep, max_chars // 2, max_chars)
        if idx != -1:
            cut = idx + 1
            break
    if cut == -1:
        cut = max_chars
    return [text[:cut].strip(), text[cut:].strip()]


def _skill_name_prof(skill) -> tuple:
    """从 skill 对象/dict 取出 (name, proficiency)。"""
    if isinstance(skill, dict):
        return skill.get("name", "") or "", skill.get("proficiency", "") or ""
    return getattr(skill, "name", "") or "", getattr(skill, "proficiency", "") or ""


def render(user, jd_summary, personal_summary, self_evaluation, out_dir: Path,
           target_position="", sections=None, fabricated=None,
           skills_curated=None, expected_salary="", include_salary=True,
           length="fresh", custom_pages=None) -> str:
    """填充工程师模板并返回生成的 docx 路径。"""
    def enabled(key: str) -> bool:
        return _enabled(sections, key)

    budget = get_length_budget(length, custom_pages)
    projects, internships, works, others = _merge_sections(
        user, fabricated, limit=budget["max_experiences"])

    doc = Document(str(TEMPLATE_PATH))
    paras = doc.paragraphs

    def setp(idx, text):
        if 0 <= idx < len(paras):
            _set_para_text(paras[idx], text)

    def clearp(idx):
        if 0 <= idx < len(paras):
            _clear(paras[idx])

    # ===== 头部 =====
    setp(SLOT_NAME, user.name or "")
    pos = target_position or user.target_position or jd_summary.get("position") or ""
    setp(SLOT_TARGET, "求职目标：" + pos if pos else "")
    clearp(SLOT_AGE)  # 模板有「年龄」字段，系统无对应数据源
    addr_parts = []
    if user.location:
        addr_parts.append("地址：" + user.location)
    if user.phone:
        addr_parts.append("手机：" + user.phone)
    setp(SLOT_ADDR_PHONE, "  ".join(addr_parts))
    setp(SLOT_EMAIL, "邮箱：" + (user.email or ""))

    # ===== 自我评价 / 个人简介 =====
    if enabled("summary") or enabled("self_evaluation"):
        summary = personal_summary if enabled("summary") else ""
        evaluation = self_evaluation if enabled("self_evaluation") else ""
        if summary and evaluation:
            setp(SLOT_SELF_1, summary)
            setp(SLOT_SELF_2, evaluation)
        else:
            text = summary or evaluation
            lines = _split_text(text, 70)
            setp(SLOT_SELF_1, lines[0])
            setp(SLOT_SELF_2, lines[1])
    else:
        clearp(SLOT_SELF_1)
        clearp(SLOT_SELF_2)

    # ===== 掌握技能（模板固定 2 条）=====
    skill_items = []
    if skills_curated:
        skill_items = [_skill_name_prof(s) for s in skills_curated]
    elif getattr(user, "skills", None):
        skill_items = [_skill_name_prof(s) for s in user.skills]
    if enabled("skills") and skill_items:
        name1, prof1 = skill_items[0]
        setp(SLOT_SKILL_1, name1 + ("（" + prof1 + "）" if prof1 else ""))
        if len(skill_items) > 1:
            name2, prof2 = skill_items[1]
            setp(SLOT_SKILL_2, name2 + ("（" + prof2 + "）" if prof2 else ""))
        else:
            clearp(SLOT_SKILL_2)
    else:
        clearp(SLOT_SKILL_1)
        clearp(SLOT_SKILL_2)

    # ===== 项目经验（模板固定 1 个完整槽位 + 散落的 2 个描述行）=====
    if enabled("projects") and projects:
        proj = projects[0]
        setp(SLOT_PROJ_TITLE, proj.title or "")
        time = f"{proj.start_date} - {proj.end_date}".strip(" -")
        role = getattr(proj, "role", "") or ""
        setp(SLOT_PROJ_TIME_ROLE, f"{time}\t{role}" if role else time)
        desc_lines = _split_text(proj.description or "", 60)
        setp(SLOT_PROJ_DESC_1, desc_lines[0])
        setp(SLOT_PROJ_DESC_2, desc_lines[1])
    else:
        clearp(SLOT_PROJ_TITLE)
        clearp(SLOT_PROJ_TIME_ROLE)
        clearp(SLOT_PROJ_DESC_1)
        clearp(SLOT_PROJ_DESC_2)

    # 外语栏：系统无对应字段，清空
    clearp(SLOT_FOREIGN)

    # ===== 教育背景 =====
    if enabled("education") and getattr(user, "educations", None):
        edu = user.educations[0]
        setp(SLOT_EDU_SCHOOL, edu.school or "")
        time = f"{edu.start_date} - {edu.end_date}".strip(" -")
        setp(SLOT_EDU_TIME, time)
        major_degree = " ".join(x for x in (edu.major, f"（{edu.degree}）" if edu.degree else "") if x)
        setp(SLOT_EDU_MAJOR, major_degree)
        course_lines = _split_text(edu.courses or "", 55)
        setp(SLOT_EDU_COURSE_1, course_lines[0])
        # 课程可能很长，模板有 3 行槽位
        remaining = course_lines[1]
        if remaining:
            sub = _split_text(remaining, 55)
            setp(SLOT_EDU_COURSE_2, sub[0])
            setp(SLOT_EDU_COURSE_3, sub[1])
        else:
            clearp(SLOT_EDU_COURSE_2)
            clearp(SLOT_EDU_COURSE_3)
    else:
        clearp(SLOT_EDU_SCHOOL)
        clearp(SLOT_EDU_TIME)
        clearp(SLOT_EDU_MAJOR)
        clearp(SLOT_EDU_COURSE_1)
        clearp(SLOT_EDU_COURSE_2)
        clearp(SLOT_EDU_COURSE_3)

    # ===== 工作经历（模板 3 个条目槽位）=====
    # 合并 works + internships 作为工作经历展示
    experiences = list(works) + list(internships)
    work_slots = [
        # (company_idx, position_idx, desc_idx, time_idx)
        (SLOT_WORK1_COMPANY, SLOT_WORK1_POSITION, SLOT_WORK1_DESC, SLOT_WORK1_TIME),
        (SLOT_WORK2_TITLE, SLOT_WORK2_POSITION, SLOT_WORK2_DESC, SLOT_WORK2_TIME),
        (SLOT_PROJ3_TITLE, SLOT_WORK3_POSITION, SLOT_WORK3_DESC, SLOT_WORK3_TIME),
    ]
    if enabled("work") or enabled("internships"):
        for i, slot in enumerate(work_slots):
            c_idx, p_idx, d_idx, t_idx = slot
            if i < len(experiences):
                exp = experiences[i]
                company = getattr(exp, "company", "") or getattr(exp, "title", "") or ""
                position = getattr(exp, "position", "") or getattr(exp, "role", "") or ""
                desc = getattr(exp, "description", "") or ""
                time = f"{exp.start_date} - {exp.end_date}".strip(" -")
                setp(c_idx, company)
                setp(p_idx, position)
                setp(d_idx, desc)
                setp(t_idx, time)
            else:
                clearp(c_idx)
                clearp(p_idx)
                clearp(d_idx)
                clearp(t_idx)
    else:
        for slot in work_slots:
            for idx in slot:
                clearp(idx)

    # 第 3 个工作条目还有一个额外描述行 + 第 2 个项目描述行
    if len(experiences) > 2 and getattr(experiences[2], "description", ""):
        extra = _split_text(experiences[2].description, 60)
        setp(SLOT_WORK3_DESC2, extra[1] if extra[1] else "")
    else:
        clearp(SLOT_WORK3_DESC2)

    if len(projects) > 1 and getattr(projects[1], "description", ""):
        setp(SLOT_PROJ2_DESC, projects[1].description)
    else:
        clearp(SLOT_PROJ2_DESC)

    # ===== 特长优势（5 条槽位，用技能摘要填充）=====
    spec_slots = [SLOT_SPEC_1, SLOT_SPEC_2, SLOT_SPEC_3, SLOT_SPEC_4, SLOT_SPEC_5]
    spec_texts = []
    if enabled("skills") and skill_items:
        for name, prof in skill_items:
            spec_texts.append(f"{name}（{prof}）" if prof else name)
    for i, idx in enumerate(spec_slots):
        setp(idx, spec_texts[i] if i < len(spec_texts) else "")

    # ===== 兴趣爱好（2 条槽位）=====
    hobbies = getattr(user, "hobbies", "") or ""
    if hobbies:
        # 按逗号/顿号/空格拆分
        import re
        parts = [x.strip() for x in re.split(r"[,，、\s]+", hobbies) if x.strip()]
        if len(parts) <= 3:
            setp(SLOT_HOBBY_1, "\t".join(parts))
            clearp(SLOT_HOBBY_2)
        else:
            setp(SLOT_HOBBY_1, "\t".join(parts[:3]))
            setp(SLOT_HOBBY_2, "\t".join(parts[3:6]))
    else:
        clearp(SLOT_HOBBY_1)
        clearp(SLOT_HOBBY_2)

    return _save_docx(doc, user, out_dir, target_position)


# 供 resume_generator.py 登记
TEMPLATE = {
    "engineer": {"label": "工程师三栏", "render": render},
}
