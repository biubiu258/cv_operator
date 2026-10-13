"""iOS 工程师风格简历模板。

源自「jianli-001 iOS工程师简历模板」，版式特征：金色系主标题、中英双语小标题、
求职目标突出展示。渲染签名与版式约定遵循 docs/模板开发规范.md。
"""
from pathlib import Path

from docx.enum.text import WD_TAB_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from ..resume_generator import (
    _add_para,
    _add_rich_runs,
    _border,
    _bullet,
    _enabled,
    _merge_sections,
    _new_doc,
    _rule,
    _save_docx,
    _set_run,
    _skill_row,
    _tech_line,
)
from ..prompting import get_length_budget

# 版式主色：金色（源自源模板 #AA894A），正文深灰 #454545
GOLD = RGBColor(0xAA, 0x89, 0x4A)
INK = RGBColor(0x45, 0x45, 0x45)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)


def _bilingual_heading(container, zh, en=""):
    """中英双语小标题：金色方块 + 中文加粗 + 英文小字 + 浅金下边框。"""
    p = _add_para(container)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("■ ")
    _set_run(r, size=8.5, bold=True, color=GOLD)
    r = p.add_run(zh)
    _set_run(r, size=11.5, bold=True, color=INK)
    if en:
        r = p.add_run("  " + en)
        _set_run(r, size=9, color=MUTED)
    _border(p, color="D8C9A8", sz="6", side="bottom", space="2")
    return p


def _gold_entry(container, left, right, tab_cm=17.3):
    """金色左侧标题 + 右侧时间右对齐的一行。"""
    p = _add_para(container)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(tab_cm), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left)
    _set_run(r, size=10.5, bold=True, color=GOLD)
    if right:
        r2 = p.add_run("\t" + right)
        _set_run(r2, size=8.5, color=MUTED)
    return p


def render(user, jd_summary, personal_summary, self_evaluation, out_dir: Path,
           target_position="", sections=None, fabricated=None,
           skills_curated=None, expected_salary="", include_salary=True,
           length="fresh", custom_pages=None) -> str:
    def enabled(key: str) -> bool:
        return _enabled(sections, key)

    budget = get_length_budget(length, custom_pages)
    projects, internships, works, others = _merge_sections(
        user, fabricated, limit=budget["max_experiences"])
    doc = _new_doc()

    # ===== 头部 =====
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(user.name or "姓名")
    _set_run(r, size=22, bold=True, color=GOLD)

    pos = target_position or user.target_position or jd_summary.get("position") or "软件工程师"
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("求职目标  ")
    _set_run(r, size=10, bold=True, color=MUTED)
    r2 = p.add_run(pos)
    _set_run(r2, size=11.5, bold=True, color=GOLD)

    if include_salary and expected_salary:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run("期望薪资  ")
        _set_run(r, size=10, bold=True, color=MUTED)
        r2 = p.add_run(expected_salary)
        _set_run(r2, size=10.5, bold=True, color=GOLD)

    items = [("电话", user.phone), ("邮箱", user.email), ("城市", user.location)]
    items = [(k, v) for k, v in items if v]
    if items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        for i, (k, v) in enumerate(items):
            if i:
                r = p.add_run("    ")
                _set_run(r, size=9, color=MUTED)
            r = p.add_run(k + " ")
            _set_run(r, size=9, color=MUTED)
            r2 = p.add_run(v)
            _set_run(r2, size=9.5, color=INK)
    _rule(doc, color="AA894A", sz="12")

    # 个人简介
    if enabled("summary") and personal_summary:
        _bilingual_heading(doc, "个人简介", "SUMMARY")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, personal_summary, size=9.5, color=INK)

    # 教育背景
    if enabled("education") and user.educations:
        _bilingual_heading(doc, "教育背景", "EDUCATION")
        for edu in user.educations:
            title = " | ".join(x for x in (edu.school, edu.major, edu.degree) if x)
            time = f"{edu.start_date} - {edu.end_date}".strip(" -")
            _gold_entry(doc, title, time)
            if edu.courses:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run("主修课程  ")
                _set_run(r, size=9, bold=True, color=MUTED)
                _add_rich_runs(p, edu.courses, size=9, color=INK)
            extra = " · ".join(x for x in (f"GPA {edu.gpa}" if edu.gpa else "",
                                           f"排名 {edu.rank}" if edu.rank else "") if x)
            if extra:
                _bullet(doc, extra, size=9)

    # 专业技能：进度条展示熟练度，空间允许时至多两个条目一行
    if enabled("skills") and (skills_curated or user.skills):
        _bilingual_heading(doc, "专业技能", "SKILLS")
        if skills_curated:
            items = [
                (s.get("name", "") if isinstance(s, dict) else s.name,
                 s.get("proficiency", "") if isinstance(s, dict) else s.proficiency)
                for s in skills_curated
            ]
            _skill_row(doc, items)
        else:
            grouped = {}
            for s in user.skills:
                grouped.setdefault(s.category or "其他", []).append(s)
            for cat, sks in grouped.items():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(cat + "  ")
                _set_run(r, size=10, bold=True, color=GOLD)
                _skill_row(doc, [(s.name, s.proficiency) for s in sks])

    # 实习经历
    if enabled("internships") and internships:
        _bilingual_heading(doc, "实习经历", "INTERNSHIP")
        for it in internships:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _gold_entry(doc, title, time)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 工作经历（杜撰）
    if enabled("work") and works:
        _bilingual_heading(doc, "工作经历", "WORK EXPERIENCE")
        for it in works:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _gold_entry(doc, title, time)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 其他经历（杜撰）
    if enabled("other") and others:
        _bilingual_heading(doc, "其他经历", "OTHER")
        for it in others:
            title = " | ".join(x for x in (it.company, it.position) if x)
            time = f"{it.start_date} - {it.end_date}".strip(" -")
            _gold_entry(doc, title, time)
            if it.description:
                _bullet(doc, it.description)
            if it.achievement:
                _bullet(doc, it.achievement)

    # 项目经历
    if enabled("projects") and projects:
        _bilingual_heading(doc, "项目经历", "PROJECT EXPERIENCE")
        for proj in projects:
            time = f"{proj.start_date} - {proj.end_date}".strip(" -")
            _gold_entry(doc, proj.title, time)
            if proj.tech_stack:
                _tech_line(doc, proj.tech_stack)
            if proj.description:
                _bullet(doc, proj.description)
            if proj.result:
                _bullet(doc, proj.result)

    # 证书 / 获奖
    if enabled("certificates") and user.certificates:
        _bilingual_heading(doc, "证书 / 获奖", "CERTIFICATES")
        for c in user.certificates:
            text = c.name + (f"（{c.date}）" if c.date else "")
            _bullet(doc, text)

    # 自我评价（置于简历最下方）
    if enabled("self_evaluation") and self_evaluation:
        _bilingual_heading(doc, "自我评价", "SELF-EVALUATION")
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _add_rich_runs(p, self_evaluation, size=9.5, color=INK)

    return _save_docx(doc, user, out_dir, target_position)


# 供 resume_generator.py 登记：key 为模板唯一标识，label 用于前端展示
TEMPLATE = {
    "ios_engineer": {"label": "iOS工程师", "render": render},
}